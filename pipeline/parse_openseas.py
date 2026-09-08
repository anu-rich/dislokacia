"""Разбор файла «Own Fleet - analysed+ Transportation summary» (KMTF UK, Ольга Васюра) — рейсы по открытым морям с 2012 г.
python3 parse_openseas.py <out.csv> <xlsx...>
Берём только объёмы и маршруты (без выручки и ставок — файл публикуется на GitHub Pages).
Файл накопительный: самый свежий файл заменяет всё.
"""
import sys, os, csv, re, datetime
import openpyxl

FIELDS = ["idx", "charterer", "ship", "type", "fleet", "loaded", "delivered", "bl", "year", "month", "load_port", "disch_port", "region", "distance", "src"]

def s(v): return "" if v is None else re.sub(r"\s+", " ", str(v)).strip()
def num(v):
    if isinstance(v, (int, float)): return float(v)
    try: return float(s(v).replace(",", "."))
    except Exception: return None
def dt(v):
    if isinstance(v, datetime.datetime): return v.strftime("%Y-%m-%d")
    t = s(v)
    for fmt in ("%d %B %Y", "%d.%m.%Y", "%Y-%m-%d"):
        try: return datetime.datetime.strptime(t, fmt).strftime("%Y-%m-%d")
        except Exception: pass
    return t

def file_date(name):
    m = re.search(r"(\d{4})[ ._-](\d{2})[ ._-](\d{2})", name)
    return f"{m.group(1)}-{m.group(2)}-{m.group(3)}" if m else ""

def parse_file(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["OwnFleetKMTF"] if "OwnFleetKMTF" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    hdr = None
    out = []
    for r in rows:
        cells = [s(c).lower() for c in r[:30]]
        if hdr is None:
            if "ship" in cells and "loaded" in cells:
                hdr = {c: i for i, c in enumerate(cells)}
            continue
        g = lambda k: r[hdr[k]] if k in hdr and hdr[k] < len(r) else None
        ship = s(g("ship"))
        if not ship: continue
        loaded = num(g("loaded"))
        if loaded is None: continue
        y = s(g("year")); mth = s(g("month"))
        try: y = int(float(y)); mth = int(float(mth))
        except Exception: continue
        out.append({"idx": s(g("internal index")), "charterer": s(g("charterer")), "ship": ship.title() if ship.isupper() and len(ship) < 8 else ship,
                    "type": s(g("type")), "fleet": s(g("fleet")), "loaded": round(loaded, 3),
                    "delivered": num(g("delivered at reported date")) or "", "bl": dt(g("bl date")), "year": y, "month": mth,
                    "load_port": s(g("loading")), "disch_port": s(g("discharging")), "region": s(g("route_region")), "distance": s(g("route distance")),
                    "src": file_date(os.path.basename(path))})
    return out

def main():
    outp = sys.argv[1]; files = sorted(sys.argv[2:], key=lambda p: file_date(os.path.basename(p)))
    if not files: return
    latest = files[-1]
    cur_src = ""
    if os.path.exists(outp):
        try: cur_src = next(csv.DictReader(open(outp, encoding="utf-8")))["src"]
        except Exception: cur_src = ""
    if file_date(os.path.basename(latest)) < cur_src:
        print(f"openseas: файл {latest} старее текущих данных ({cur_src}), пропуск", file=sys.stderr); return
    rows = parse_file(latest)
    with open(outp, "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=FIELDS); w.writeheader(); w.writerows(rows)
    print(f"openseas: {len(rows)} рейсов из {os.path.basename(latest)}", file=sys.stderr)

if __name__ == "__main__": main()
