import os
import openpyxl, glob, datetime, re, json, csv, sys
from collections import OrderedDict


def norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()

def find_header(ws, must):
    for r in range(1, 6):
        vals = [norm(c.value) for c in ws[r]]
        if any(must in v for v in vals):
            return r, vals
    return None, None

def colmap(vals, spec):
    m = {}
    for key, pats in spec.items():
        for i, v in enumerate(vals):
            if v and any(p in v for p in pats) and key not in m:
                # avoid double match e.g. "начало выгрузки" vs "начало погрузки"
                m[key] = i
    return m

TANKER_SPEC = OrderedDict([
    ("vessel", ["наименование судна"]),
    ("shpr", ["shpr"]),
    ("op", ["операция"]),
    ("arr", ["прихода на рейд"]),
    ("berth_at", ["постановки к причалу"]),
    ("berth", ["причал №", "пр.№"]),
    ("load_start", ["начало погрузки"]),
    ("load_end", ["окончание погрузки"]),
    ("draft", ["осадка"]),
    ("cargo", ["кол-во взятого груза", "количество груза", "кол-во груза"]),
    ("dep", ["отхода"]),
])
BULK_SPEC = OrderedDict([
    ("vessel", ["наименование судна"]),
    ("op", ["операция"]),
    ("arr", ["приход на рейд", "прихода на рейд"]),
    ("berth_at", ["постановка к причалу", "постановки к причалу"]),
    ("berth", ["пр.№", "причал №"]),
    ("unload_start", ["начало выгрузки"]),
    ("unload_end", ["окончание выгрузки"]),
    ("dep", ["отхода"]),
])

def to_dt(v):
    if isinstance(v, datetime.datetime): return v
    if isinstance(v, datetime.date): return datetime.datetime(v.year, v.month, v.day)
    return None

def to_num(v):
    if v is None: return None
    if isinstance(v, (int, float)): return float(v)
    s = str(v).replace(" ", "").replace("\xa0", "").replace(",", ".")
    m = re.search(r"\d+(\.\d+)?", s)
    return float(m.group()) if m else None

def clean_name(s):
    s = re.sub(r"\s+", " ", str(s or "")).strip()
    fixes = {"GC Barys": "GC Barys", "GC Berkut": "GC Berkut", "GC Sunkar": "GC Sunkar",
             "Пр. Гейдар Алиев": "Пр.Гейдар Алиев", "Президент Гейдар Алиев": "Пр.Гейдар Алиев",
             "ТК АКТАУ": "ТК Актау", "Тк Актау": "ТК Актау"}
    return fixes.get(s, s)

def parse_file(f, tankers, bulk, bunker):
    wb = openpyxl.load_workbook(f, data_only=True, read_only=True)
    src = f.split("/")[-1][:16]
    # ---- tankers
    ws = wb["Статистика по танкерам"]
    hr, vals = find_header(ws, "наименование судна")
    cm = colmap(vals, TANKER_SPEC)
    term_col = cm["shpr"] + 1  # unlabeled column after SHPR = terminal
    year = None
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        a = row[0]
        if isinstance(a, str) and re.match(r"^\s*20\d\d", a): year = int(a.strip()[:4]); continue
        arr = to_dt(row[cm["arr"]]); dep = to_dt(row[cm["dep"]])
        if not a or not (arr or dep): continue
        rec = dict(vessel=clean_name(a), shpr=norm(row[cm["shpr"]]).upper() if row[cm["shpr"]] else "",
                   terminal=str(row[term_col] or "").strip(), berth=str(row[cm["berth"]] or "").strip(),
                   arr=arr, berth_at=to_dt(row[cm["berth_at"]]), load_start=to_dt(row[cm["load_start"]]),
                   load_end=to_dt(row[cm["load_end"]]), draft=to_num(row[cm["draft"]]), cargo=to_num(row[cm["cargo"]]),
                   dep=dep, src=src)
        key = (rec["vessel"], (dep or arr).strftime("%Y-%m-%d %H:%M"))
        tankers[key] = rec  # later files override
    # ---- bulk
    ws = wb["Статистика сухогр."]
    hr, vals = find_header(ws, "наименование судна")
    cm = colmap(vals, BULK_SPEC)
    # cargo columns: two blocks (unload / load): вид груза, осадка, кол-во, дфэ
    idx_kind = [i for i, v in enumerate(vals) if "вид груза" in v]
    idx_qty = [i for i, v in enumerate(vals) if "кол-во груза" in v]
    idx_teu = [i for i, v in enumerate(vals) if "дфэ" in v]
    idx_draft = [i for i, v in enumerate(vals) if "осадка" in v]
    idx_ls = [i for i, v in enumerate(vals) if "начало погрузки" in v]
    idx_le = [i for i, v in enumerate(vals) if "окончание погрузки" in v]
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        a = row[0]
        arr = to_dt(row[cm["arr"]]); dep = to_dt(row[cm["dep"]])
        if not a or not isinstance(a, str) or not (arr or dep): continue
        g = lambda idxs, k: (row[idxs[k]] if len(idxs) > k and idxs[k] < len(row) else None)
        rec = dict(vessel=clean_name(a), berth=str(row[cm["berth"]] or "").strip(), arr=arr, berth_at=to_dt(row[cm["berth_at"]]),
                   unload_start=to_dt(row[cm["unload_start"]]), unload_end=to_dt(row[cm["unload_end"]]),
                   in_kind=norm(g(idx_kind, 0)), in_qty=to_num(g(idx_qty, 0)), in_qty_raw=str(g(idx_qty, 0) or "").strip(), in_teu=to_num(g(idx_teu, 0)), in_draft=to_num(g(idx_draft, 0)),
                   load_start=to_dt(g(idx_ls, 0)), load_end=to_dt(g(idx_le, 0)),
                   out_kind=norm(g(idx_kind, 1)), out_qty=to_num(g(idx_qty, 1)), out_qty_raw=str(g(idx_qty, 1) or "").strip(), out_teu=to_num(g(idx_teu, 1)), out_draft=to_num(g(idx_draft, 1)),
                   dep=dep, src=src)
        key = (rec["vessel"], (dep or arr).strftime("%Y-%m-%d %H:%M"))
        bulk[key] = rec
    # ---- bunker: два блока (п.Актау / п.Баку №146 район и др.), ищем заголовки «Наименование судна»
    wsn = [n for n in wb.sheetnames if n.startswith("Статистика бункеровки")][0]
    ws = wb[wsn]
    G = [list(r) for r in ws.iter_rows(values_only=True)]
    hr = next((r for r in range(min(6, len(G))) if any(isinstance(v, str) and norm(v).startswith("наименование судна") for v in G[r])), None)
    if hr is not None:
        cols = [c for c, v in enumerate(G[hr]) if isinstance(v, str) and norm(v).startswith("наименование судна")]
        titles = {}
        for c in cols:
            t = ""
            for r in range(hr - 1, max(-1, hr - 3), -1):
                for cc in range(c, min(c + 9, len(G[r]))):
                    if isinstance(G[r][cc], str) and norm(G[r][cc]).startswith("п."): t = norm(G[r][cc]); break
                if t: break
            titles[c] = t
        for row in G[hr + 1:]:
            for c in cols:
                g = lambda k: row[c + k] if c + k < len(row) else None
                a = g(0)
                if not a or not isinstance(a, str): continue
                arr = to_dt(g(1)); bs = to_dt(g(3)); be = to_dt(g(4))
                if not (bs or arr): continue
                place = str(g(8) or "").strip() or titles.get(c, "") or "Актау"
                rec = dict(vessel=clean_name(a), arr=arr, b_start=bs, b_end=be, dt=to_num(g(5)), tt=to_num(g(6)), dep=to_dt(g(7)), port=place, src=src)
                key = (rec["vessel"], (bs or arr).strftime("%Y-%m-%d %H:%M"))
                bunker[key] = rec
    print(os.path.basename(f), len(tankers), len(bulk), len(bunker), file=sys.stderr)

def dump(d, path):
    rows = sorted(d.values(), key=lambda r: (r.get("dep") or r.get("arr") or r.get("b_start")))
    if not rows: return
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        w.writeheader()
        for r in rows:
            w.writerow({k: (v.strftime("%Y-%m-%d %H:%M") if isinstance(v, datetime.datetime) else ("" if v is None else v)) for k, v in r.items()})

if __name__ == "__main__":
    tankers, bulk, bunker = {}, {}, {}
    for f in sorted(sys.argv[1:]):
        parse_file(f, tankers, bulk, bunker)
    dump(tankers, "data/tankers.csv"); dump(bulk, "data/bulk.csv"); dump(bunker, "data/bunker.csv")
